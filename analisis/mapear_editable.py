"""
Separa el texto EDITABLE del texto generado por Citavi.

Todo w:t que descienda de un w:sdt cuyo campo sea CitaviPlaceholder es
resultado de campo: si se edita, Word lo revierte al actualizar las citas.
Ese texto NO se debe tocar.
"""
import docx, re, collections
from docx.oxml.ns import qn

DOCX = "/projects/sandbox/hola/Tesis Jaime Fredy Horacio Avance 18 - comillas corregidas.docx"
doc = docx.Document(DOCX)
SDT = qn('w:sdt')

def es_de_citavi(nodo):
    """True si el w:t esta dentro de un control con campo Citavi."""
    n = nodo.getparent()
    while n is not None:
        if n.tag == SDT:
            if b"Citavi" in __import__("lxml.etree", fromlist=["etree"]).tostring(n):
                return True
        n = n.getparent()
    return False

PATRONES = {
    "espacio antes de coma/punto": re.compile(r"\s+[,;.](?:\s|$)"),
    "dos o mas espacios":          re.compile(r"[^\s]  +[^\s]"),
    "falta espacio tras punto":    re.compile(r"[a-záéíóúñ]\.[A-ZÁÉÍÓÚÑ]"),
    "coma sin espacio detras":     re.compile(r",[a-zA-ZáéíóúñÁÉÍÓÚÑ]"),
}

editable = collections.Counter()
en_citavi = collections.Counter()
ejemplos = collections.defaultdict(list)

for i, p in enumerate(doc.paragraphs):
    for t in p._p.iter(qn('w:t')):
        if not t.text:
            continue
        destino = en_citavi if es_de_citavi(t) else editable
        for nombre, rx in PATRONES.items():
            hits = len(rx.findall(t.text))
            if hits:
                destino[nombre] += hits
                if destino is editable and len(ejemplos[nombre]) < 5:
                    ejemplos[nombre].append((i, t.text[:90]))

print(f"{'PATRON':<32}{'EDITABLE':>10}{'DENTRO DE CITAVI':>19}")
print("-" * 62)
for nombre in PATRONES:
    print(f"{nombre:<32}{editable[nombre]:>10}{en_citavi[nombre]:>19}")

print("\nEjemplos de lo que SI se puede corregir:")
for nombre, ejs in ejemplos.items():
    print(f"\n  {nombre}:")
    for i, e in ejs:
        print(f"    [{i}] {e!r}")

# Runs con fuente distinta de Calibri
print("\n" + "=" * 62)
print("RUNS CON FUENTE EXPLICITA DISTINTA DE CALIBRI")
print("=" * 62)
for i, p in enumerate(doc.paragraphs):
    for r in p._p.iter(qn('w:r')):
        rpr = r.find(qn('w:rPr'))
        if rpr is None:
            continue
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            continue
        vals = {rf.get(qn(a)) for a in ('w:ascii', 'w:hAnsi', 'w:cs') if rf.get(qn(a))}
        malas = {v for v in vals if v and "Calibri" not in v}
        if malas:
            txt = "".join(t.text or "" for t in r.iter(qn('w:t')))
            print(f"  [parrafo {i}] fuente={malas} texto={txt[:70]!r}")

# ¿tablas con fuentes distintas?
print("\nRUNS NO-CALIBRI DENTRO DE TABLAS:")
cnt = collections.Counter()
for tbl in doc.tables:
    for r in tbl._tbl.iter(qn('w:r')):
        rpr = r.find(qn('w:rPr'))
        rf = rpr.find(qn('w:rFonts')) if rpr is not None else None
        if rf is None:
            continue
        for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
            v = rf.get(qn(a))
            if v and "Calibri" not in v:
                cnt[v] += 1
print("  ", dict(cnt) or "ninguno")
