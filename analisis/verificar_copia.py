"""Compara el original con la copia corregida: solo deben cambiar las comillas."""
import zipfile, docx, re, hashlib
from docx.oxml.ns import qn

A = "/projects/sandbox/hola/Tesis Jaime Fredy Horacio Avance 18 (1).docx"
B = "/projects/sandbox/hola/Tesis Jaime Fredy Horacio Avance 18 - comillas corregidas.docx"

za, zb = zipfile.ZipFile(A), zipfile.ZipFile(B)
na, nb = set(za.namelist()), set(zb.namelist())
print("=== ESTRUCTURA DEL PAQUETE ===")
print(f"  partes en el original: {len(na)} | en la copia: {len(nb)}")
print(f"  partes faltantes:  {na - nb or 'ninguna'}")
print(f"  partes añadidas:   {nb - na or 'ninguna'}")
distintas = [n for n in sorted(na & nb)
             if hashlib.md5(za.read(n)).hexdigest() != hashlib.md5(zb.read(n)).hexdigest()]
print(f"  partes con contenido distinto: {distintas}")

da, db = docx.Document(A), docx.Document(B)

def real(p):
    return "".join(t.text or "" for t in p._p.iter(qn('w:t')))

MAY, MIN = "A-ZÁÉÍÓÚÑ", "a-záéíóúñ"
CITA = re.compile(rf"\([^()\n]{{0,80}}?[{MAY}][^()\n]{{1,80}}?,?\s*(?:19|20)\d{{2}}[a-z]?"
                  rf"(?:\s*[-–]\s*\d{{2,4}})?\s*(?:,\s*p+\.?\s*[\d\-]+)?\)")
CITA_N = re.compile(rf"[{MAY}][\w{MIN}\.]*(?:\s+(?:y|and|&|et\s+al\.?|de|del|la)?\s*[{MAY}][\w{MIN}\.]*){{0,4}}"
                    rf"\s*\(\s*(?:19|20)\d{{2}}[a-z]?\s*\)")

def metricas(d, z):
    body = d.element.body
    txt = "\n".join(real(p) for p in d.paragraphs)
    citas = set()
    for p in d.paragraphs:
        t = real(p)
        for rx in (CITA, CITA_N):
            for m in rx.finditer(t):
                citas.add(m.group(0).strip())
    return {
        "parrafos": len(d.paragraphs),
        "tablas": len(d.tables),
        "campos Citavi": d.element.body.xml.count("CitaviPlaceholder"),
        "controles w:sdt": d.element.body.xml.count("<w:sdt>"),
        "imagenes en media/": sum(1 for n in z.namelist() if n.startswith("word/media/")),
        "citas detectadas": len(citas),
        "palabras": len(txt.split()),
        "caracteres ¨": txt.count("\u00a8"),
        "comillas “": txt.count("\u201c"),
        "comillas ”": txt.count("\u201d"),
    }

ma, mb = metricas(da, za), metricas(db, zb)
print("\n=== CONTENIDO ===")
print(f"{'métrica':<22}{'original':>12}{'copia':>10}   ")
print("-" * 50)
for k in ma:
    igual = "=" if ma[k] == mb[k] else "  <-- CAMBIO"
    print(f"{k:<22}{ma[k]:>12}{mb[k]:>10}   {igual}")

# El texto debe ser identico salvo los caracteres de comilla
ta = "\n".join(real(p) for p in da.paragraphs)
tb = "\n".join(real(p) for p in db.paragraphs)
norm = lambda s: s.replace("\u00a8", "@").replace("\u201c", "@").replace("\u201d", "@")
norm2 = lambda s: re.sub(r"\s+", " ", norm(s))
print(f"\nTexto idéntico ignorando comillas y espacios: {norm2(ta) == norm2(tb)}")

# Que la copia abra sin errores y las citas sean las mismas
cita_a = {m.group(0).strip() for p in da.paragraphs for m in CITA.finditer(real(p))}
cita_b = {m.group(0).strip() for p in db.paragraphs for m in CITA.finditer(real(p))}
print(f"Mismas citas en ambos archivos:               {cita_a == cita_b}")
print(f"Citas perdidas:                               {cita_a - cita_b or 'ninguna'}")
