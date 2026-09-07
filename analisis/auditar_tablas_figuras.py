"""Audita la numeracion de tablas y figuras: secuencia, huecos, duplicados,
concordancia con los indices y con las llamadas del texto."""
import docx, re, collections
from docx.oxml.ns import qn

DOCX = "/projects/sandbox/hola/Tesis Jaime Fredy Horacio Avance 18 - REVISADA.docx"
doc = docx.Document(DOCX)
ps = doc.paragraphs

def real(p):
    return "".join(t.text or "" for t in p._p.iter(qn('w:t')))

# ---------- rotulos reales en el cuerpo (estilo APA: numero y titulo aparte) ----------
rot_tabla, rot_fig = [], []
for i, p in enumerate(ps):
    t = real(p).strip()
    st = p.style.name
    m = re.match(r"^Tabla\s+(\d+)\s*\.?\s*(.*)$", t, re.I)
    if m and i > 196:            # despues de los indices
        titulo = m.group(2).strip()
        if not titulo:            # estilo APA: el titulo va en el parrafo siguiente
            j = i + 1
            titulo = real(ps[j]).strip() if j < len(ps) else ""
        rot_tabla.append((i, int(m.group(1)), titulo, st))
    m = re.match(r"^Figura\s+(\d+)\s*\.?\s*(.*)$", t, re.I)
    if m and i > 196:
        titulo = m.group(2).strip()
        if not titulo:
            j = i + 1
            titulo = real(ps[j]).strip() if j < len(ps) else ""
        rot_fig.append((i, int(m.group(1)), titulo, st))

def informe(nombre, rots, n_objetos=None):
    print("=" * 84)
    print(f"{nombre}: {len(rots)} rotulos encontrados en el cuerpo")
    print("=" * 84)
    nums = [n for _, n, _, _ in rots]
    dup = [n for n, c in collections.Counter(nums).items() if c > 1]
    print(f"  secuencia: {nums}")
    print(f"  duplicados: {sorted(dup) or 'ninguno'}")
    if nums:
        esperado = set(range(1, max(nums) + 1))
        print(f"  huecos:     {sorted(esperado - set(nums)) or 'ninguno'}")
        desorden = [(nums[k - 1], nums[k]) for k in range(1, len(nums)) if nums[k] < nums[k - 1]]
        print(f"  fuera de orden: {desorden or 'ninguno'}")
    print()
    for i, n, tit, st in rots:
        print(f"    [{i:4d}] {nombre[:-1]} {n:2d}  {tit[:66]}")
    print()

informe("TABLAS", rot_tabla)
informe("FIGURAS", rot_fig)

# ---------- objetos reales ----------
print("=" * 84)
print("OBJETOS REALES EN EL DOCUMENTO")
print("=" * 84)
print(f"  tablas (w:tbl):            {len(doc.tables)}")
n_img = doc.element.body.xml.count("<a:blip")
print(f"  imagenes incrustadas:      {n_img}")
print(f"  rotulos 'Tabla N':         {len(rot_tabla)}")
print(f"  rotulos 'Figura N':        {len(rot_fig)}")

# ---------- indices ----------
print("\n" + "=" * 84)
print("INDICES (parrafos 101-196) frente a los rotulos del cuerpo")
print("=" * 84)
idx_t, idx_f = [], []
for i in range(101, 197):
    t = real(ps[i]).strip()
    m = re.match(r"^Tabla\s+(\d+)", t, re.I)
    if m: idx_t.append(int(m.group(1)))
    m = re.match(r"^Figura\s+(\d+)", t, re.I)
    if m: idx_f.append(int(m.group(1)))
print(f"  indice de tablas:  {len(idx_t)} entradas -> {idx_t}")
print(f"  cuerpo:            {len(rot_tabla)} rotulos -> {[n for _,n,_,_ in rot_tabla]}")
print(f"  faltan en el indice: {sorted(set(n for _,n,_,_ in rot_tabla) - set(idx_t)) or 'ninguna'}")
print(f"  sobran en el indice: {sorted(set(idx_t) - set(n for _,n,_,_ in rot_tabla)) or 'ninguna'}")
print()
print(f"  indice de figuras: {len(idx_f)} entradas -> {idx_f}")
print(f"  cuerpo:            {len(rot_fig)} rotulos -> {[n for _,n,_,_ in rot_fig]}")
print(f"  faltan en el indice: {sorted(set(n for _,n,_,_ in rot_fig) - set(idx_f)) or 'ninguna'}")
print(f"  sobran en el indice: {sorted(set(idx_f) - set(n for _,n,_,_ in rot_fig)) or 'ninguna'}")

# ---------- llamadas en el texto ----------
print("\n" + "=" * 84)
print("LLAMADAS EN EL TEXTO ('la Tabla N', 'la Figura N') que no existen")
print("=" * 84)
existe_t = {n for _, n, _, _ in rot_tabla}
existe_f = {n for _, n, _, _ in rot_fig}
malas = 0
for i in range(197, len(ps)):
    t = real(ps[i])
    for m in re.finditer(r"\b[Tt]abla\s+(\d+)", t):
        n = int(m.group(1))
        if n not in existe_t:
            print(f"  [{i}] menciona Tabla {n} y no hay rotulo: …{t[max(0,m.start()-45):m.end()+35]}…")
            malas += 1
    for m in re.finditer(r"\b[Ff]igura\s+(\d+)", t):
        n = int(m.group(1))
        if n not in existe_f:
            print(f"  [{i}] menciona Figura {n} y no hay rotulo: …{t[max(0,m.start()-45):m.end()+35]}…")
            malas += 1
if not malas:
    print("  ninguna")

# ---------- notas de fuente ----------
print("\n" + "=" * 84)
print("ROTULOS SIN NOTA DE FUENTE debajo")
print("=" * 84)
faltan = []
for i, n, tit, st in rot_tabla + rot_fig:
    ventana = " ".join(real(ps[j]) for j in range(i, min(i + 6, len(ps))))
    if not re.search(r"Nota\.|Fuente:|Elaboración propia|Adaptado de|Tomado de", ventana, re.I):
        faltan.append((i, n, tit[:55]))
print(f"  {len(faltan)} sin nota localizable:")
for i, n, tit in faltan:
    print(f"    [{i}] {n}: {tit}")
