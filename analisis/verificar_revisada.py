"""
Verifica que la version REVISADA solo difiera en espacios, fuente y niveles de
titulo. Comprueba palabra por palabra que el contenido es identico.
"""
import zipfile, docx, re, hashlib, collections
from docx.oxml.ns import qn

A = "/projects/sandbox/hola/Tesis Jaime Fredy Horacio Avance 18 (1).docx"          # original intacto
B = "/projects/sandbox/hola/Tesis Jaime Fredy Horacio Avance 18 - REVISADA.docx"   # entregable

za, zb = zipfile.ZipFile(A), zipfile.ZipFile(B)
print("=" * 72)
print("1. INTEGRIDAD DEL PAQUETE")
print("=" * 72)
na, nb = set(za.namelist()), set(zb.namelist())
print(f"   partes: {len(na)} -> {len(nb)}   faltantes: {na-nb or 'ninguna'}   nuevas: {nb-na or 'ninguna'}")
dif = [n for n in sorted(na & nb)
       if hashlib.md5(za.read(n)).hexdigest() != hashlib.md5(zb.read(n)).hexdigest()]
print(f"   partes modificadas: {dif}")

da, db = docx.Document(A), docx.Document(B)
def real(p): return "".join(t.text or "" for t in p._p.iter(qn('w:t')))

print("\n" + "=" * 72)
print("2. ESTRUCTURA")
print("=" * 72)
for nombre, fa, fb in [
    ("parrafos",            len(da.paragraphs), len(db.paragraphs)),
    ("tablas",              len(da.tables), len(db.tables)),
    ("imagenes",            sum(1 for n in za.namelist() if n.startswith("word/media/")),
                            sum(1 for n in zb.namelist() if n.startswith("word/media/"))),
    ("campos Citavi",       da.element.body.xml.count("CitaviPlaceholder"),
                            db.element.body.xml.count("CitaviPlaceholder")),
    ("controles w:sdt",     da.element.body.xml.count("<w:sdt>"),
                            db.element.body.xml.count("<w:sdt>")),
    ("saltos de pagina",    da.element.body.xml.count('w:type="page"'),
                            db.element.body.xml.count('w:type="page"')),
    ("notas al pie",        da.element.body.xml.count("footnoteReference"),
                            db.element.body.xml.count("footnoteReference")),
]:
    print(f"   {nombre:<20}{fa:>8} -> {fb:>8}   {'OK' if fa==fb else '*** REVISAR ***'}")

print("\n" + "=" * 72)
print("3. CONTENIDO: comparacion palabra por palabra")
print("=" * 72)
def palabras(d):
    out = []
    for p in d.paragraphs:
        out += real(p).split()
    return out
wa, wb = palabras(da), palabras(db)
print(f"   palabras: {len(wa)} -> {len(wb)}")

# normalizar solo las comillas (se cambiaron ¨ por “ ”) y comparar
Q = lambda s: s.replace("\u00a8", "\u201c").replace("\u201d", "\u201c")
na_, nb_ = [Q(w) for w in wa], [Q(w) for w in wb]
if na_ == nb_:
    print("   >>> IDENTICAS. No se cambio, añadio ni elimino ninguna palabra.")
else:
    print("   >>> DIFERENCIAS:")
    import difflib
    for linea in list(difflib.unified_diff(na_, nb_, lineterm="", n=1))[:60]:
        print("      ", linea)

print("\n" + "=" * 72)
print("4. FUENTE")
print("=" * 72)
def fuentes(d):
    c = collections.Counter()
    for rpr in d.element.body.iter(qn('w:rPr')):
        rf = rpr.find(qn('w:rFonts'))
        if rf is not None:
            for at in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
                v = rf.get(qn(at))
                if v: c[v] += 1
    return c
print("   original:", dict(fuentes(da)))
print("   revisada:", dict(fuentes(db)))
z = zipfile.ZipFile(B)
th = z.read("word/theme/theme1.xml").decode("utf8")
print("   tema  ->  minor:", re.findall(r'<a:minorFont>.*?typeface="([^"]*)"', th, re.S)[:1],
      " major:", re.findall(r'<a:majorFont>.*?typeface="([^"]*)"', th, re.S)[:1])
dd = re.search(r"<w:docDefaults>.*?</w:docDefaults>", z.read("word/styles.xml").decode("utf8"), re.S)
print("   fuente por defecto:", set(re.findall(r'w:(?:ascii|hAnsi)="([^"]+)"', dd.group(0))))
noCal = [v for v in fuentes(db) if "Calibri" not in v]
print(f"   >>> runs con fuente distinta de Calibri: {noCal or 'NINGUNO'}")

print("\n" + "=" * 72)
print("5. TIPOGRAFIA: incidencias restantes")
print("=" * 72)
PAT = {
    "espacio antes de coma/;": re.compile(r"[ ]+[,;]"),
    "espacio antes de punto":  re.compile(r"[ ]+\.(?=\s|$)"),
    "coma sin espacio":        re.compile(r",(?=[A-Za-zÁÉÍÓÚÑáéíóúñ])"),
    "espacios dobles":         re.compile(r"(?<=\S)[ ]{2,}(?=\S)"),
    "caracter ¨":              re.compile(r"\u00a8"),
}
for nombre, rx in PAT.items():
    ca = sum(len(rx.findall(real(p))) for p in da.paragraphs)
    cb = sum(len(rx.findall(real(p))) for p in db.paragraphs)
    print(f"   {nombre:<26}{ca:>5} -> {cb:>5}")

print("\n" + "=" * 72)
print("6. JERARQUIA DE TITULOS")
print("=" * 72)
for nombre, d in (("original", da), ("revisada", db)):
    saltos, prev = [], None
    for i, p in enumerate(d.paragraphs):
        m = re.match(r"(?:Heading|Título|Titulo)\s*(\d)", p.style.name)
        if not m: continue
        lvl = int(m.group(1))
        if prev and lvl > prev + 1:
            saltos.append((i, real(p).strip()[:45]))
        prev = lvl
    print(f"   {nombre}: {len(saltos)} saltos incoherentes {saltos if saltos else ''}")
