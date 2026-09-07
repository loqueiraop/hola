import docx, collections, re, zipfile
from docx.oxml.ns import qn

DOCX = "/projects/sandbox/hola/Tesis Jaime Fredy Horacio Avance 18 - comillas corregidas.docx"
doc = docx.Document(DOCX)
W = qn  # alias

# ---------- 1. FUENTES ----------
print("=" * 78)
print("FUENTES USADAS (rFonts explicitos en runs)")
print("=" * 78)
fuentes = collections.Counter()
tam = collections.Counter()
for p in doc.paragraphs:
    for r in p._r_lst if hasattr(p, "_r_lst") else p.runs:
        pass
for p in doc.element.body.iter(qn('w:r')):
    rf = p.find(qn('w:rPr') + '/' + qn('w:rFonts')) if False else None
for rpr in doc.element.body.iter(qn('w:rPr')):
    rf = rpr.find(qn('w:rFonts'))
    if rf is not None:
        for at in ('w:ascii', 'w:hAnsi', 'w:cs'):
            v = rf.get(qn(at))
            if v:
                fuentes[v] += 1
    sz = rpr.find(qn('w:sz'))
    if sz is not None:
        tam[sz.get(qn('w:val'))] += 1
for f, n in fuentes.most_common():
    print(f"  {n:6d}  {f}")
print("\nTAMAÑOS (media-puntos -> pt):")
for s, n in tam.most_common(12):
    print(f"  {n:6d}  {s} = {int(s)/2:g} pt")

# ---------- 2. FUENTE POR DEFECTO ----------
z = zipfile.ZipFile(DOCX)
styles = z.read("word/styles.xml").decode("utf8")
m = re.search(r"<w:docDefaults>.*?</w:docDefaults>", styles, re.S)
print("\nFUENTE POR DEFECTO del documento (docDefaults):")
if m:
    for at in re.findall(r'w:(?:ascii|hAnsi|cs)="([^"]+)"', m.group(0)):
        print("   ", at)
    for at in re.findall(r'<w:sz w:val="(\d+)"', m.group(0)):
        print("    tamaño:", int(at) / 2, "pt")

# estilo Normal
m2 = re.search(r'<w:style [^>]*w:styleId="Normal".*?</w:style>', styles, re.S)
print("\nEstilo 'Normal':")
if m2:
    print("   ", re.findall(r'w:(?:ascii|hAnsi)="([^"]+)"', m2.group(0)) or "(hereda docDefaults)")

# ---------- 3. ESTILOS DE TITULO: coherencia de niveles ----------
print("\n" + "=" * 78)
print("JERARQUIA DE TITULOS: saltos de nivel incoherentes")
print("=" * 78)
ant = None
for i, p in enumerate(doc.paragraphs):
    st = p.style.name
    if not re.match(r"(Heading|Título|Titulo)\s*\d", st):
        continue
    lvl = int(re.sub(r"\D", "", st))
    txt = "".join(t.text or "" for t in p._p.iter(qn('w:t'))).strip()
    if ant is not None and lvl > ant + 1:
        print(f"  [{i}] salta de nivel {ant} a {lvl}: {txt[:60]}")
    ant = lvl

# ---------- 4. TIPOGRAFIA SUCIA ----------
print("\n" + "=" * 78)
print("INCONSISTENCIAS TIPOGRAFICAS")
print("=" * 78)
def real(p):
    return "".join(t.text or "" for t in p._p.iter(qn('w:t')))

pat = {
    "espacio antes de coma/punto":      re.compile(r"\s+[,;.](?:\s|$)"),
    "dos o mas espacios seguidos":      re.compile(r"[^\s]  +[^\s]"),
    "espacio antes de cierre )":        re.compile(r"\s+\)"),
    "falta espacio tras punto":         re.compile(r"[a-záéíóúñ]\.[A-ZÁÉÍÓÚÑ]"),
    "coma sin espacio detras":          re.compile(r",[a-zA-ZáéíóúñÁÉÍÓÚÑ]"),
}
tot = collections.Counter()
ejemplos = collections.defaultdict(list)
for i, p in enumerate(doc.paragraphs):
    t = real(p)
    for nombre, rx in pat.items():
        for m in rx.finditer(t):
            tot[nombre] += 1
            if len(ejemplos[nombre]) < 4:
                ejemplos[nombre].append((i, t[max(0, m.start() - 40):m.end() + 40]))
for nombre in pat:
    print(f"\n  {nombre}: {tot[nombre]}")
    for i, e in ejemplos[nombre]:
        print(f"     [{i}] …{e}…")
