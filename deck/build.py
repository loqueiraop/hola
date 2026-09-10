"""
Genera "Contrato de franquicia — Acuerdos Comerciales, Grupo 1".
Presentación de 15 minutos, 16:9.

Estructura: un componente por diapositiva, según lo pedido.
  1 Portada
  2 Definición conceptual
  3 Cómo se aplica
  4 Partes involucradas
  5 Ventajas y desventajas
  6 Intro al video entrevista
  7 Bibliografía
  8 Cierre
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from design import (
    SW, SH, M, CW,
    INK, INK_SOFT, MUTED, BG, CARD, CARD_ALT, LINE, ACCENT, ACCENT_TN,
    ON_DARK, ON_DARK_SF, SERIF, SANS,
    textbox, para, rich, rect, rule, vrule, canvas, header, footer, notes,
    card, bullets,
)

DARK_CARD = RGBColor(0x22, 0x33, 0x4C)
HAIR = RGBColor(0x2C, 0x3E, 0x59)

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)

_n = 0


def page(dark=False, numbered=True):
    global _n
    _n += 1
    s = canvas(prs, dark=dark)
    if numbered:
        footer(s, _n, dark=dark)
    return s


def block_label(slide, x, y, w, text, *, dark=False, size=8.5):
    tf = textbox(slide, x, y, w, 0.20)
    para(tf, text, first=True, size=size, color=ACCENT, bold=True, caps=True,
         spc=150, line=1.0)


# =====================================================================
# 1 · PORTADA
# =====================================================================
s = page(dark=True, numbered=False)

rect(s, 8.40, 1.55, 3.95, 3.95, fill=None, line_color=HAIR, line_w=1.1,
     shape=MSO_SHAPE.OVAL)
rect(s, 8.02, 4.62, 0.76, 0.76, fill=ACCENT, shape=MSO_SHAPE.OVAL)
vrule(s, 7.55, 1.55, 3.95, color=HAIR)

tf = textbox(s, M, 1.52, 6.4, 0.24)
para(tf, "Maestría · Acuerdos Comerciales · Grupo 1", first=True, size=10,
     color=ACCENT, bold=True, caps=True, spc=200, line=1.0)

tf = textbox(s, M, 1.96, 6.6, 1.70)
para(tf, "Contrato de", first=True, size=48, font=SERIF, color=ON_DARK, line=1.02)
para(tf, "franquicia", size=48, font=SERIF, color=ON_DARK, line=1.02)

rule(s, M, 3.76, 0.9)

tf = textbox(s, M, 4.02, 6.1, 0.9)
para(tf, "Una figura de colaboración empresarial que organiza la expansión de "
         "una marca sin renunciar al control de su estándar.",
     first=True, size=13.5, color=ON_DARK_SF, italic=True, line=1.35)

rect(s, M, 5.52, 6.1, 0.008, fill=HAIR)

tf = textbox(s, M, 5.74, 6.6, 0.28)
para(tf, "María Isabel Acosta · Xilena Blanco · María Daniela García · "
         "Fredy Valladares", first=True, size=10.5, color=ON_DARK, line=1.2)

tf = textbox(s, M, 6.10, 6.6, 0.24)
para(tf, "Definición · aplicación · partes · ventajas y desventajas · entrevista",
     first=True, size=9, color=ON_DARK_SF, caps=True, spc=110, line=1.0)

tf = textbox(s, SW - M - 3.0, 6.10, 3.0, 0.24)
para(tf, "Septiembre 2026", first=True, size=9, color=ON_DARK_SF,
     align=PP_ALIGN.RIGHT, caps=True, spc=110, line=1.0)

notes(s, "0:00–0:30 · Saludo, tema y anuncio de que la exposición cierra con una "
         "entrevista a un directivo que opera más de 40 franquicias.")


# =====================================================================
# 2 · DEFINICIÓN CONCEPTUAL
# =====================================================================
s = page()
header(s, "Componente 1", "Definición conceptual",
       subtitle="Dos nociones que conviene no confundir: el modelo de negocio y "
                "el instrumento que lo ordena.")

CW2 = (CW - 0.33) / 2
CH = 3.40
IN = 0.36

# --- La franquicia
x = M
card(s, x, 2.20, CW2, CH, fill=CARD, accent_top=True)
tf = textbox(s, x + IN, 2.46, CW2 - 2 * IN, 0.30)
para(tf, "La franquicia", first=True, size=17, font=SERIF, color=INK, line=1.0)
tf = textbox(s, x + IN, 2.84, CW2 - 2 * IN, 0.72)
para(tf, "Modelo de negocio mediante el cual una empresa autoriza a otra a "
         "utilizar su marca, sus conocimientos, sus procedimientos y su sistema "
         "de operación para desarrollar una actividad comercial.",
     first=True, size=11, color=INK_SOFT, line=1.32)
block_label(s, x + IN, 3.64, CW2 - 2 * IN, "Lo que se transfiere")
bullets(s, x + IN, 3.88, CW2 - 2 * IN, 0.80, [
    "La marca: nombre comercial, logotipo e imagen",
    "El know-how: recetas, procesos y manuales validados",
    "El sistema: diseño, atención, publicidad y estándares",
], size=10, gap=5)
tf = textbox(s, x + IN, 4.76, CW2 - 2 * IN, 0.44)
para(tf, "En la práctica: distintos dueños, misma experiencia para el consumidor.",
     first=True, size=10, color=INK, italic=True, line=1.25)
tf = textbox(s, x + IN, 5.26, CW2 - 2 * IN, 0.24)
para(tf, "(Speicher Mendiola, 2024)", first=True, size=9, color=MUTED, line=1.1)

# --- El contrato
x = M + CW2 + 0.33
card(s, x, 2.20, CW2, CH, fill=CARD_ALT, border=LINE)
rect(s, x, 2.20, CW2, 0.045, fill=INK)
tf = textbox(s, x + IN, 2.46, CW2 - 2 * IN, 0.30)
para(tf, "El contrato de franquicia", first=True, size=17, font=SERIF, color=INK,
     line=1.0)
tf = textbox(s, x + IN, 2.84, CW2 - 2 * IN, 0.72)
para(tf, "Instrumento que fija las condiciones de la relación entre las partes: "
         "derechos, obligaciones, responsabilidades y términos bajo los cuales "
         "se desarrollará la franquicia.",
     first=True, size=11, color=INK_SOFT, line=1.32)
block_label(s, x + IN, 3.64, CW2 - 2 * IN, "Rasgos que lo definen")
bullets(s, x + IN, 3.88, CW2 - 2 * IN, 1.34, [
    "Bilateral y oneroso: obligaciones recíprocas y precio",
    "De colaboración, no de simple intercambio",
    "De tracto sucesivo: se ejecuta de forma continuada",
    "Contrato marco: regula la pertenencia a una red",
    "Necesariamente incompleto, para poder adaptarse",
], size=10, gap=5)
tf = textbox(s, x + IN, 5.26, CW2 - 2 * IN, 0.24)
para(tf, "(Apolo Apolo et al., 2022; Kim y Tiwana, 2022)", first=True, size=9,
     color=MUTED, line=1.1)

# --- banda: reconocimiento jurídico
card(s, M, 5.78, CW, 0.86, fill=INK, border=None)
tf = textbox(s, M + 0.4, 5.94, CW - 0.8, 0.22)
para(tf, "Un reconocimiento jurídico desigual", first=True, size=9.5,
     color=ACCENT, bold=True, caps=True, spc=150, line=1.0)
tf = textbox(s, M + 0.4, 6.18, CW - 0.8, 0.42)
para(tf, "Argentina lo tipificó (arts. 1512–1524 CCyC); España lo regula por vía "
         "sectorial (Ley 7/1996); Ecuador lo reconoce en su Código de Comercio; "
         "Colombia y Perú lo mantienen como contrato atípico, de modo que el peso "
         "recae en la redacción del contrato y en la autonomía de la voluntad.",
     first=True, size=10, color=ON_DARK, line=1.3)

notes(s, "0:30–3:15 · Distinguir franquicia (modelo de negocio) de contrato "
         "(instrumento jurídico). Cerrar con la banda: donde no hay tipificación, "
         "el contrato hace de norma. Ese es el punto que sostiene toda la "
         "exposición.")


# =====================================================================
# 3 · CÓMO SE APLICA
# =====================================================================
s = page()
header(s, "Componente 2", "¿Cómo se aplica?",
       subtitle="El ciclo que va del acuerdo entre dos partes a la operación "
                "cotidiana de una red.")

etapas = [
    ("01", "Selección del franquiciado",
     "La elección del inversionista adecuado es determinante para el éxito del "
     "sistema (Calderón-Monge et al., 2021)."),
    ("02", "Formalización del acuerdo",
     "Se firma el contrato y quedan fijadas las condiciones económicas, "
     "operativas y legales de la relación."),
    ("03", "Transferencia y capacitación",
     "Entrenamiento del propietario y del personal, manuales de operación, "
     "asesoría administrativa y apoyo en marketing."),
    ("04", "Operación del establecimiento",
     "El franquiciado administra su negocio, pero respetando los estándares "
     "definidos por el franquiciador."),
    ("05", "Seguimiento y supervisión",
     "Control continuo para que todas las unidades sostengan los estándares que "
     "identifican a la marca."),
]
cwe = (CW - 4 * 0.19) / 5
for i, (num, t, d) in enumerate(etapas):
    x = M + i * (cwe + 0.19)
    card(s, x, 2.26, cwe, 2.90, accent_top=True)
    tf = textbox(s, x + 0.26, 2.54, cwe - 0.52, 0.5)
    para(tf, num, first=True, size=26, font=SERIF, color=ACCENT, line=0.95)
    tf = textbox(s, x + 0.26, 3.12, cwe - 0.52, 0.56)
    para(tf, t, first=True, size=11.5, font=SERIF, color=INK, line=1.14)
    tf = textbox(s, x + 0.26, 3.78, cwe - 0.52, 1.3)
    para(tf, d, first=True, size=9.5, color=INK_SOFT, line=1.28)
    if i < 4:
        rect(s, x + cwe + 0.055, 3.82, 0.08, 0.08, fill=LINE,
             shape=MSO_SHAPE.OVAL)

# --- dos cierres: el deber previo y la fórmula del modelo
x = M
card(s, x, 5.30, CW2, 1.28, fill=CARD_ALT, border=LINE)
block_label(s, x + 0.28, 5.48, CW2 - 0.56, "Antes de firmar", size=9)
tf = textbox(s, x + 0.28, 5.72, CW2 - 0.56, 0.82)
para(tf, "El franquiciador debe entregar información sobre su identidad, el "
         "sector, las características de la franquicia y la organización de la "
         "red al menos 20 días hábiles antes de la firma. Omitirla o deformarla "
         "genera responsabilidad aunque el contrato nunca se firme (Conde Gómez, "
         "2020).",
     first=True, size=9.5, color=INK_SOFT, line=1.26)

x = M + CW2 + 0.33
card(s, x, 5.30, CW2, 1.28, fill=INK, border=None)
rect(s, x, 5.30, CW2, 0.045, fill=ACCENT)
block_label(s, x + 0.28, 5.48, CW2 - 0.56, "La fórmula del modelo", size=9)
tf = textbox(s, x + 0.28, 5.72, CW2 - 0.56, 0.82)
para(tf, "La franquicia combina independencia empresarial con un nivel de control "
         "establecido por el titular de la marca. El franquiciado administra su "
         "propio negocio, pero no decide sobre todo: opera dentro de un estándar "
         "que no le pertenece.",
     first=True, size=9.5, color=ON_DARK, line=1.26)

notes(s, "3:15–6:00 · Recorrer las cinco etapas con ritmo. Detenerse en la 1 (la "
         "selección explica por qué unas redes funcionan y otras no) y en el "
         "recuadro del deber previo: la responsabilidad puede nacer antes de que "
         "exista contrato.")


# =====================================================================
# 4 · PARTES INVOLUCRADAS
# =====================================================================
s = page()
header(s, "Componente 3", "Partes involucradas",
       subtitle="Dos empresarios jurídicamente independientes que, sin embargo, "
                "se necesitan mutuamente.")

partes = [
    ("El franquiciador", "también franquiciante u otorgante", True,
     "Marca y signos distintivos · know-how probado · manuales operativos · "
     "asistencia técnica y comercial",
     "Cobrar canon y regalías · exigir confidencialidad del know-how · "
     "actualizar los manuales · exigir la estandarización de la red",
     "Comunicar el know-how · prestar asistencia técnica · entregar la "
     "información precontractual · responder ante litigios de propiedad "
     "intelectual"),
    ("El franquiciado", "o franquiciatario", False,
     "Canon de entrada e inversión inicial · regalías periódicas · gestión y "
     "equipo humano · conocimiento del mercado local",
     "Capacitación y asistencia permanente · manual operativo completo · "
     "exclusividad territorial si se pactó · adaptar el producto al mercado "
     "local",
     "Pagar canon y regalías · mantener la confidencialidad · seguir el manual "
     "operativo · permitir auditorías · no competir con la red"),
]

for i, (t, alias, primary, aporta, derechos, oblig) in enumerate(partes):
    x = M + i * (CW2 + 0.33)
    y = 2.14
    card(s, x, y, CW2, 3.34, fill=(CARD if primary else CARD_ALT),
         border=LINE, accent_top=primary)
    if not primary:
        rect(s, x, y, CW2, 0.045, fill=INK)
    tf = textbox(s, x + IN, y + 0.24, CW2 - 2 * IN, 0.30)
    para(tf, t, first=True, size=17, font=SERIF, color=INK, line=1.0)
    tf = textbox(s, x + IN, y + 0.58, CW2 - 2 * IN, 0.18)
    para(tf, alias, first=True, size=8.5, color=MUTED, italic=True, caps=True,
         spc=110, line=1.0)
    for k, (lab, txt) in enumerate((("Qué aporta", aporta),
                                    ("Sus derechos", derechos),
                                    ("Sus obligaciones", oblig))):
        by = y + 0.82 + k * 0.84
        block_label(s, x + IN, by, CW2 - 2 * IN, lab)
        tf = textbox(s, x + IN, by + 0.20, CW2 - 2 * IN, 0.52)
        para(tf, txt, first=True, size=9.5, color=INK_SOFT, line=1.26)

card(s, M, 5.62, CW, 0.52, fill=CARD_ALT, border=None)
tf = textbox(s, M + 0.34, 5.70, CW - 0.68, 0.40)
rich(tf, [("Y un tercero cuando la red se internacionaliza.  ",
           {"size": 9.5, "color": INK, "bold": True}),
          ("Aparece el subfranquiciante o máster franquiciado, que sin ser "
           "titular de la marca asume frente a terceros las funciones propias de "
           "un franquiciador (Speicher Mendiola, 2024; Conde Gómez, 2020).",
           {"size": 9.5, "color": INK_SOFT})],
     first=True, line=1.26)

card(s, M, 6.20, CW, 0.50, fill=INK, border=None)
tf = textbox(s, M + 0.34, 6.26, CW - 0.68, 0.40)
rich(tf, [("Independencia declarada, pero no absoluta.  ",
           {"size": 9.5, "color": ACCENT, "bold": True}),
          ("Argentina afirma que las partes son independientes (art. 1520 CCyC) "
           "y el proyecto colombiano descarta la solidaridad pasiva salvo pacto "
           "expreso. Aun así, la jurisprudencia española la extiende al "
           "franquiciador cuando el daño deriva de sus instrucciones concretas.",
           {"size": 9.5, "color": ON_DARK})],
     first=True, line=1.26)

notes(s, "6:00–9:00 · Presentar a las dos partes. Insistir en que el franquiciado "
         "es empresario, no empleado. No leer las listas completas: destacar dos "
         "o tres de cada bloque. Cerrar con la banda inferior, que es el punto de "
         "discusión jurídica más interesante.")


# =====================================================================
# 5 · VENTAJAS Y DESVENTAJAS
# =====================================================================
s = page()
header(s, "Componente 4", "Principales ventajas y desventajas",
       subtitle="Estandarizar es lo que crea el valor de la red y, a la vez, lo "
                "que restringe al franquiciado.")

ventajas = [
    ("Reconocimiento de marca",
     "se arranca con una marca posicionada: capta clientes antes y gana "
     "reputación más rápido."),
    ("Modelo de negocio validado",
     "se accede a un negocio ya probado y se reduce la incertidumbre de crear "
     "una marca desde cero."),
    ("Menor riesgo empresarial",
     "registra tasas de supervivencia superiores a las de los negocios "
     "independientes."),
    ("Capacitación y asistencia",
     "formación técnica inicial y acompañamiento continuo en marketing, "
     "operaciones y gestión."),
    ("Apoyo en marketing",
     "las campañas y las redes sociales se diseñan y coordinan desde la marca."),
    ("Economías de escala",
     "la red negocia como bloque: mejores condiciones y menores costos de "
     "adquisición."),
]
desventajas = [
    ("Costos elevados",
     "inversión inicial alta más pagos continuos por regalías y publicidad."),
    ("Limitación de la autonomía",
     "debe seguir con rigor las políticas, estándares y procedimientos del "
     "franquiciador."),
    ("Regalías al margen del resultado",
     "se pagan aunque el negocio no genere las ganancias esperadas."),
    ("Conflictos entre las partes",
     "diferencias sobre estrategias, proveedores autorizados o condiciones "
     "económicas."),
    ("Cláusulas estrictas",
     "exclusividad, duración, renovación y restricciones que limitan la "
     "flexibilidad."),
    ("Dependencia de la reputación ajena",
     "un golpe a la imagen del franquiciador repercute sobre toda la red."),
]

for i, (titulo, items, positivo) in enumerate(
        (("Ventajas", ventajas, True), ("Desventajas", desventajas, False))):
    x = M + i * (CW2 + 0.33)
    card(s, x, 2.34, CW2, 4.24, fill=(CARD if positivo else CARD_ALT),
         border=LINE, accent_top=positivo)
    if not positivo:
        rect(s, x, 2.34, CW2, 0.045, fill=INK)
    tf = textbox(s, x + IN, 2.60, CW2 - 2 * IN, 0.30)
    para(tf, titulo, first=True, size=17, font=SERIF, color=INK, line=1.0)
    tf = textbox(s, x + IN, 3.02, CW2 - 2 * IN, 3.4)
    for k, (t, d) in enumerate(items):
        rich(tf, [(f"{k + 1}. ", {"size": 9.5, "color": ACCENT, "bold": True}),
                  (t + ": ", {"size": 9.5, "color": INK, "bold": True}),
                  (d, {"size": 9.5, "color": INK_SOFT})],
             first=(k == 0), space_after=11, line=1.26)

notes(s, "9:00–11:15 · No leer las doce. Agrupar: las ventajas se resumen en que "
         "se compra tiempo y se comparte infraestructura; las desventajas, en que "
         "se paga con dinero y con autonomía. Cerrar con la idea del subtítulo: "
         "son dos caras de la misma decisión.")


# =====================================================================
# 6 · INTRO AL VIDEO ENTREVISTA
# =====================================================================
s = page(dark=True)
header(s, "Componente 5", "Video entrevista", dark=True,
       subtitle="Entrevista sobre las claves del éxito del Grupo David en la "
                "expansión de franquicias.")

# --- invitado
x = M
card(s, x, 2.30, CW2, 2.10, fill=DARK_CARD, border=None)
rect(s, x, 2.30, CW2, 0.045, fill=ACCENT)
block_label(s, x + IN, 2.54, CW2 - 2 * IN, "Invitado", size=9)
tf = textbox(s, x + IN, 2.82, CW2 - 2 * IN, 0.36)
para(tf, "Gerardo Marcano", first=True, size=22, font=SERIF, color=ON_DARK,
     line=1.0)
tf = textbox(s, x + IN, 3.24, CW2 - 2 * IN, 0.24)
para(tf, "Vicepresidente de operaciones · Grupo David", first=True, size=10.5,
     color=ACCENT, bold=True, line=1.15)
tf = textbox(s, x + IN, 3.54, CW2 - 2 * IN, 0.70)
para(tf, "Veinte años en la organización. Ha sido gerente de tiendas minoristas, "
         "gerente de operaciones y director de retail. Hoy supervisa la operación "
         "integral del grupo, que incluye más de 40 franquicias.",
     first=True, size=10, color=ON_DARK_SF, line=1.32)

# --- escala de la operación
x = M + CW2 + 0.33
card(s, x, 2.30, CW2, 2.10, fill=DARK_CARD, border=None)
rect(s, x, 2.30, CW2, 0.045, fill=HAIR)
block_label(s, x + IN, 2.54, CW2 - 2 * IN, "Escala de la operación", size=9)
datos = [("23+", "marcas operadas"), ("350+", "tiendas en la región"),
         ("40+", "franquicias gestionadas"), ("20+", "plataformas de e-commerce")]
dw = (CW2 - 2 * IN - 0.24) / 2
for k, (fig, lab) in enumerate(datos):
    col, row = k % 2, k // 2
    dx = x + IN + col * (dw + 0.24)
    dy = 2.86 + row * 0.78
    tf = textbox(s, dx, dy, dw, 0.32)
    para(tf, fig, first=True, size=19, font=SERIF, color=ON_DARK, line=0.95)
    tf = textbox(s, dx, dy + 0.34, dw, 0.32)
    para(tf, lab, first=True, size=8.5, color=ON_DARK_SF, caps=True, spc=90,
         line=1.15)

# --- qué observar
block_label(s, M, 4.68, CW, "Qué observar mientras vemos el video", size=9.5)
observar = [
    ("01", "Cómo el reconocimiento de marca decide si una apertura se hace o no."),
    ("02", "Hasta dónde puede adaptarse el modelo sin romper el estándar: la "
           "tropicalización."),
    ("03", "El manual operativo en la práctica: marco flexible o camisa de "
           "fuerza."),
]
cw3 = (CW - 2 * 0.26) / 3
for k, (num, t) in enumerate(observar):
    x = M + k * (cw3 + 0.26)
    card(s, x, 5.00, cw3, 1.18, fill=DARK_CARD, border=None)
    rect(s, x, 5.00, 0.045, 1.18, fill=ACCENT)
    tf = textbox(s, x + 0.30, 5.22, cw3 - 0.60, 0.28)
    para(tf, num, first=True, size=13, font=SERIF, color=ACCENT, line=1.0)
    tf = textbox(s, x + 0.30, 5.54, cw3 - 0.60, 0.50)
    para(tf, t, first=True, size=10.5, color=ON_DARK, line=1.26)

tf = textbox(s, M, 6.34, CW, 0.24)
para(tf, "Ejes de la conversación: estrategias de expansión · innovación y "
         "omnicanalidad · cultura de replicación operativa",
     first=True, size=9, color=ON_DARK_SF, italic=True, line=1.1)

notes(s, "11:15–12:00 · Presentar al invitado y justificar por qué su testimonio "
         "importa: opera la escala completa del modelo en la región. Dar las tres "
         "claves de observación y reproducir el video. Dejar esta diapositiva en "
         "pantalla mientras se proyecta.")


# =====================================================================
# 7 · BIBLIOGRAFÍA
# =====================================================================
refs = [
    ("Apolo Apolo, B. V., Reinoso Galarza, B. A., y Cando Pacheco, J. de J. "
     "(2022). ", "Análisis crítico del contrato de franquicia en la legislación "
     "ecuatoriana, especial referencia a la responsabilidad del franquiciado. "
     "RECIMUNDO, 6(3), 459–474."),
    ("Calderón-Monge, E., Pastor-Sanz, I., y Sendra-García, J. (2021). ",
     "How to select franchisees: A model proposal. Journal of Business Research, "
     "135, 676–684."),
    ("Candela, P. (2023, 1 de febrero). ",
     "Good papers: La extensión de responsabilidad al franquiciador por los "
     "actos del franquiciatario. Almacén de Derecho."),
    ("Conde Gómez, G. (2020). ",
     "El contrato de franquicia en la Propuesta de Anteproyecto de Ley de Código "
     "Mercantil tras el Dictamen del Consejo de Estado. Revista de Derecho UNED, "
     "(26), 77–103."),
    ("Congreso de la Nación Argentina. (2014). ",
     "Ley 26.994: Código Civil y Comercial de la Nación, arts. 1512–1524, "
     "Contrato de franquicia."),
    ("EAE Business School. (2025). ",
     "Ventajas e inconvenientes del contrato de franquicia."),
    ("Federal Trade Commission. (2020). ",
     "A consumer's guide to buying a franchise."),
    ("Franchise Business Review. (2025). ",
     "What are the advantages and disadvantages of franchising?"),
    ("Kim, S. K., y Tiwana, A. (2022). ",
     "Franchising contracts as routines: Untangling the adaptive value of "
     "incomplete contracts. Journal of Business Research, 152, 177–190."),
    ("Loaiza, L. M. (2012). ",
     "Panorama del contrato de franquicia en Colombia: entre las estipulaciones "
     "contractuales, el Código de Ética y la Norma Técnica [proyecto de grado, "
     "Universidad Icesi]."),
    ("Ministerio de Comercio, Industria y Turismo. (2021). ",
     "Proyecto de decreto sobre las franquicias como alternativa para el "
     "emprendimiento y la expansión de las mipymes [no expedido]."),
    ("NerdWallet. (2025). ",
     "Franchising: Advantages, disadvantages and how to buy a franchise."),
    ("Speicher Mendiola, C. G. (2024). ",
     "La tipificación del contrato de franquicia como alternativa para impulsar "
     "el desarrollo de las empresas peruanas en el mercado local. Ius et Praxis, "
     "(58), 23–39."),
    ("Superintendencia de Industria y Comercio. (2010). ",
     "Concepto sobre el contrato de franquicia [citado en Loaiza, 2012]."),
]

s = page()
header(s, "Referencias", "Bibliografía")
colw = (CW - 0.5) / 2
for i, (autor, resto) in enumerate(refs):
    col, row = i // 7, i % 7
    x = M + col * (colw + 0.5)
    y = 2.20 + row * 0.62
    tf = textbox(s, x, y, colw, 0.58)
    rich(tf, [(autor, {"size": 8.5, "color": INK, "bold": True}),
              (resto, {"size": 8.5, "color": INK_SOFT})],
         first=True, line=1.28)
tf = textbox(s, M, 6.58, CW, 0.24)
para(tf, "Referencias en formato APA 7.ª edición.", first=True, size=8.5,
     color=MUTED, italic=True, line=1.1)

notes(s, "Diapositiva de respaldo. No se lee en la exposición.")


# =====================================================================
# 8 · CIERRE
# =====================================================================
s = page(dark=True, numbered=False)

rect(s, 8.9, 2.2, 3.2, 3.2, fill=None, line_color=HAIR, line_w=1.1,
     shape=MSO_SHAPE.OVAL)
rect(s, 10.05, 4.68, 0.62, 0.62, fill=ACCENT, shape=MSO_SHAPE.OVAL)

tf = textbox(s, M, 2.42, 7.2, 0.24)
para(tf, "Contrato de franquicia · Acuerdos Comerciales", first=True, size=10,
     color=ACCENT, bold=True, caps=True, spc=200, line=1.0)

tf = textbox(s, M, 2.80, 7.2, 0.9)
para(tf, "Gracias", first=True, size=52, font=SERIF, color=ON_DARK, line=1.0)

rule(s, M, 3.88, 0.9)

tf = textbox(s, M, 4.16, 6.4, 0.4)
para(tf, "Quedamos atentos a sus preguntas y comentarios.", first=True,
     size=13.5, color=ON_DARK_SF, italic=True, line=1.3)

tf = textbox(s, M, 5.22, 6.4, 0.24)
para(tf, "Grupo 1", first=True, size=9, color=ACCENT, bold=True, caps=True,
     spc=160, line=1.0)
tf = textbox(s, M, 5.52, 6.4, 1.0)
for i, nombre in enumerate(["María Isabel Acosta", "Xilena Blanco",
                            "María Daniela García", "Fredy Valladares"]):
    para(tf, nombre, first=(i == 0), size=11.5, color=ON_DARK, space_after=3,
         line=1.15)

notes(s, "Tras el video: retomar en una frase la idea de cierre —el éxito del "
         "modelo no se decide en el texto del contrato, sino en la disciplina con "
         "que se replica y se adapta el estándar— y abrir preguntas.")


# =====================================================================
out = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "Contrato de Franquicia - Grupo 1.pptx")
prs.save(out)
print(f"OK · {len(prs.slides._sldIdLst)} diapositivas")
print(out)
